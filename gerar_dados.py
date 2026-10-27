"""Gera a planilha de exemplo entrada/pedidos.xlsx (dados fictícios).

A maioria dos pedidos está certa, mas alguns têm os problemas que
aparecem na vida real, para mostrar como o robô lida com eles.
As datas são calculadas a partir de hoje, para o exemplo nunca vencer.
"""
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from formatar import AZUL, FONTE

hoje = date.today()


def daqui(dias):
    return hoje + timedelta(days=dias)


pedidos = [
    # Cliente, Produto, Quantidade, Entrega, Observação
    ("C001", "P01", 40, daqui(3), "Entregar pela manhã"),
    ("C002", "P07", 15, daqui(2), ""),
    ("C003", "P02", 30, daqui(5), ""),
    ("C004", "P08", 60, daqui(4), "Portão lateral"),
    ("C009", "P03", 20, daqui(3), ""),            # cliente que não existe no sistema
    ("C005", "P06", 12, daqui(7), ""),
    ("C006", "P04", 500, daqui(6), ""),           # mais do que o estoque
    ("C007", "P05", 10, daqui(2), "Ligar antes"),
    ("C008", "P04", 8, daqui(-2), ""),            # data de entrega já passou
    ("C001", "P12", 5, daqui(3), ""),             # produto que não existe
    ("C003", "P08", 25, daqui(8), ""),
    ("C002", "P02", None, daqui(4), ""),          # quantidade em branco
]

Path("entrada").mkdir(exist_ok=True)
caminho = "entrada/pedidos.xlsx"
tabela = pd.DataFrame(pedidos, columns=["Cliente", "Produto", "Quantidade",
                                        "Entrega", "Observação"])
tabela["Quantidade"] = tabela["Quantidade"].astype("Int64")
tabela.to_excel(caminho, index=False)

# Deixa a planilha com cara de planilha de verdade: cabeçalho, larguras, data
livro = load_workbook(caminho)
aba = livro.active
for celula, largura in zip(aba[1], (10, 10, 12, 13, 24)):
    celula.fill = AZUL
    celula.font = Font(name=FONTE, bold=True, color="FFFFFF")
    celula.alignment = Alignment(horizontal="center")
    aba.column_dimensions[celula.column_letter].width = largura
for (celula,) in aba.iter_rows(min_row=2, min_col=4, max_col=4):
    celula.number_format = "DD/MM/YYYY"
    celula.alignment = Alignment(horizontal="center")
aba.freeze_panes = "A2"
livro.save(caminho)
print(f"{len(tabela)} pedidos gravados em {caminho}")
