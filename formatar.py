from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

AZUL = PatternFill("solid", fgColor="1F4E78")
VERDE = PatternFill("solid", fgColor="DEF7E5")
VERMELHO = PatternFill("solid", fgColor="FDE2E2")
SEM_COR = PatternFill(fill_type=None)
FONTE = "Calibri"
COLUNAS_DO_ROBO = {"Situação": 11, "Pedido nº": 11, "Motivo": 48}


def anotar_resultados(caminho, resultados):
    """Escreve o resultado de cada linha na própria planilha de pedidos.

    Usa o openpyxl direto (e não o pandas) para mexer só nas colunas do
    robô e nas cores, sem estragar a formatação que o cliente já tinha.
    resultados = {linha_do_excel: (situacao, numero, motivo, print_da_tela)}
    """
    livro = load_workbook(caminho)
    aba = livro.active
    colunas = {celula.value: celula.column for celula in aba[1] if celula.value}

    for nome, largura in COLUNAS_DO_ROBO.items():
        if nome not in colunas:
            colunas[nome] = aba.max_column + 1
            cabecalho = aba.cell(1, colunas[nome], nome)
            cabecalho.fill = AZUL
            cabecalho.font = Font(name=FONTE, bold=True, color="FFFFFF")
            cabecalho.alignment = Alignment(horizontal="center")
            aba.column_dimensions[cabecalho.column_letter].width = largura

    for linha, (situacao, numero, motivo, imagem) in resultados.items():
        for nome, valor in (("Situação", situacao), ("Pedido nº", numero)):
            aba.cell(linha, colunas[nome], valor).alignment = Alignment(horizontal="center")
        celula_motivo = aba.cell(linha, colunas["Motivo"], motivo or None)
        celula_motivo.hyperlink = None
        celula_motivo.font = Font(name=FONTE)
        if imagem:
            # Clicar no motivo abre o print da tela no momento do erro
            celula_motivo.hyperlink = imagem
            celula_motivo.font = Font(name=FONTE, color="0563C1", underline="single")
        cor = VERDE if situacao == "Lançado" else VERMELHO
        for coluna in range(1, aba.max_column + 1):
            aba.cell(linha, coluna).fill = cor

    aba.freeze_panes = "A2"
    # Na impressão, todas as colunas cabem numa folha deitada
    aba.page_setup.orientation = "landscape"
    aba.sheet_properties.pageSetUpPr.fitToPage = True
    aba.page_setup.fitToWidth, aba.page_setup.fitToHeight = 1, 0
    livro.save(caminho)
